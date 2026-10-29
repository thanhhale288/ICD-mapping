"""
P0.3 / P0.4 — Trích Excel bàn giao ra CSV (chỉ đọc, không ghi Excel).

Đầu ra:
  data/icd_hien_trang.csv  — tổng hợp + hạng mục chi tiết
  data/icd_keys.csv        — bảng khóa ICD (tên, địa bàn, QĐ, chủ ĐT)
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / "Thông tin ICD 34 tỉnh năm 2026 - đã kiểm tra.xlsx"
OUT_HIEN_TRANG = ROOT / "data" / "icd_hien_trang.csv"
OUT_KEYS = ROOT / "data" / "icd_keys.csv"

# R7 — nhóm tên trùng / dễ nhầm (đánh dấu để duyệt tay)
NHOM_TRUNG: dict[int, str] = {
    15: "Long Bình",
    98: "Long Bình",
    126: "Long Bình",
    58: "Tân Lập",
    121: "Tân Lập",
    18: "Thạnh Phước/Thanh Phước",
    116: "Thạnh Phước/Thanh Phước",
    17: "Cái Mép",
    115: "Cái Mép",
    21: "Thanh Thủy",
    71: "Thanh Thủy",
    8: "Quế Võ/Mộc Bài",
    19: "Quế Võ/Mộc Bài",
}

# Tỉnh/thành (mới hoặc cũ) hay gặp trong file — dùng để tách địa bàn
TINH_ALIASES: dict[str, str] = {
    "tp.hcm": "TP. Hồ Chí Minh",
    "tp. hcm": "TP. Hồ Chí Minh",
    "tp hồ chí minh": "TP. Hồ Chí Minh",
    "tp. hồ chí minh": "TP. Hồ Chí Minh",
    "hồ chí minh": "TP. Hồ Chí Minh",
    "bà rịa–vũng tàu": "Bà Rịa – Vũng Tàu",
    "bà rịa-vũng tàu": "Bà Rịa – Vũng Tàu",
    "bà rịa – vũng tàu": "Bà Rịa – Vũng Tàu",
}

# Sheet 08 chưa có (Phase 5) — bổ sung khóa tối thiểu từ Tổng hợp + địa bàn biết trước
STT8_FALLBACK = {
    "ten_chinh": "Cảng cạn Tân Cảng Quế Võ",
    "ten_goi_khac": "Tân Cảng Quế Võ; Quế Võ",
    "tinh_moi": "Bắc Ninh",
    "tinh_cu": "Bắc Ninh",
    "huyen_cu": "Quế Võ",
    "vi_tri": "",
    "so_qd": "",
    "chu_dau_tu": "",
    "don_vi_khai_thac": "",
    "khoa_dia_ban": "Quế Võ; Bắc Ninh",
    "sheet": "",
    "sheet_thieu": "1",
}


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sheet_stt(name: str) -> int | None:
    m = re.match(r"^0*(\d+)", name)
    return int(m.group(1)) if m else None


def normalize_space(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def canon_tinh(name: str) -> str:
    key = normalize_space(name).lower().replace("–", "-")
    if key in TINH_ALIASES:
        return TINH_ALIASES[key]
    # giữ nguyên, chỉ chuẩn hóa khoảng trắng / gạch
    return normalize_space(name.replace("–", "–"))


def strip_prefix_ten(ten: str) -> str:
    s = ten
    s = re.sub(r"^(Cảng cạn\s*)+", "", s, flags=re.IGNORECASE)
    s = re.sub(r"^\(ICD\)\s*", "", s, flags=re.IGNORECASE)
    s = re.sub(r"^ICD\s+", "", s, flags=re.IGNORECASE)
    s = re.sub(r"\s*\(giai đoạn\s*\d+\)\s*$", "", s, flags=re.IGNORECASE)
    s = re.sub(r"\s+giai đoạn\s*\d+\s*$", "", s, flags=re.IGNORECASE)
    return normalize_space(s)


def split_paren(ten: str) -> tuple[str, list[str]]:
    """Tách phần trong ngoặc làm tên gọi khác / địa bàn phụ."""
    alts: list[str] = []
    base = ten
    for m in re.finditer(r"\(([^)]+)\)", ten):
        inner = normalize_space(m.group(1))
        if inner:
            alts.append(inner)
    # bỏ ngoặc khỏi tên gốc khi ngoặc là ghi chú địa bàn/tên khác
    base = normalize_space(re.sub(r"\s*\([^)]+\)", "", ten))
    return base, alts


def parse_vi_tri(vi_tri: str) -> tuple[str, str, str, str]:
    """
    Trả về (tinh_moi, tinh_cu, huyen_cu, khoa_dia_ban).
    Heuristic dựa trên mẫu vị trí trong file bàn giao.
    """
    text = normalize_space(vi_tri)
    if not text:
        return "", "", "", ""

    tinh_moi = ""
    tinh_cu = ""
    huyen_cu = ""
    parts: list[str] = []

    # nay thuộc X (ưu tiên TP.HCM / TP. Hồ Chí Minh trước khi cắt tại dấu chấm)
    m_nay = re.search(
        r"nay thuộc\s+(TP\.?\s*HCM|TP\.?\s*Hồ Chí Minh|[^.;,]+?)(?:\s*\(|\s*\.|$)",
        text,
        flags=re.IGNORECASE,
    )
    if m_nay:
        raw_nay = normalize_space(m_nay.group(1))
        if re.search(r"HCM|Hồ Chí Minh", raw_nay, re.I):
            tinh_moi = "TP. Hồ Chí Minh"
        else:
            tinh_moi = canon_tinh(raw_nay)

    # (X cũ) hoặc trước đây ... X
    for m in re.finditer(r"\(([^)]+?)\s+cũ\)", text, flags=re.IGNORECASE):
        tinh_cu = canon_tinh(m.group(1)) if not tinh_cu else tinh_cu
        parts.append(canon_tinh(m.group(1)) + " (cũ)")

    m_truoc = re.search(
        r"(?:trước đây|địa danh trước sắp xếp)[^.;]*?([A-ZÀ-Ỹ][^.;,]{2,40})",
        text,
        flags=re.IGNORECASE,
    )
    # Mẫu quy hoạch: "Khu vực <địa điểm>, <tỉnh> (theo tên..."
    m_khu = re.match(
        r"Khu vực\s+(.+?),\s*([^,(]+?)\s*(?:\(theo|\(|$)",
        text,
        flags=re.IGNORECASE,
    )
    if m_khu:
        dia_diem = normalize_space(m_khu.group(1))
        tinh = canon_tinh(m_khu.group(2))
        if not tinh_moi:
            tinh_moi = tinh
        # địa điểm có thể chứa "(tỉnh cũ)" hoặc huyện
        base_dd, alts_dd = split_paren(dia_diem)
        parts.append(base_dd)
        parts.extend(alts_dd)
        # ngoặc trong địa điểm: tỉnh/huyện cũ hoặc ghi chú địa bàn — không phải tên thương hiệu
        known_place = {
            "hưng yên",
            "long an",
            "an giang",
            "gò dầu",
            "cái mép",
            "tuy phước",
            "lộc ninh",
            "krông búk",
            "hậu giang",
            "tp.hcm",
            "bình dương",
            "bà rịa – vũng tàu",
        }
        for a in alts_dd:
            if re.search(r"\bcũ\b", a, re.I):
                tinh_cu = canon_tinh(re.sub(r"\s*cũ\s*", "", a, flags=re.I))
            elif a.lower() in known_place or re.search(
                r"tỉnh|huyện|thành phố|tp\.?", a, re.I
            ):
                if a.lower() == tinh.lower():
                    continue
                # Tân Lập (Long An) khi tỉnh mới là Tây Ninh → Long An là tỉnh/địa bàn cũ
                if not tinh_cu and a.lower() != (tinh_moi or "").lower():
                    tinh_cu = canon_tinh(a)
                elif not huyen_cu:
                    huyen_cu = a

    # "tỉnh X"
    for m in re.finditer(r"tỉnh\s+([^,.;()]+)", text, flags=re.IGNORECASE):
        t = canon_tinh(m.group(1))
        if not tinh_moi:
            tinh_moi = t
        parts.append(t)

    # TP / thành phố
    for m in re.finditer(
        r"(?:TP\.?|thành phố)\s*([^,.;()]+)",
        text,
        flags=re.IGNORECASE,
    ):
        t = canon_tinh("TP. " + m.group(1) if not m.group(1).lower().startswith("hồ") else "TP. Hồ Chí Minh")
        # đặc biệt HCM
        raw = normalize_space(m.group(1))
        if re.search(r"hồ chí minh|hcm", raw, re.I):
            t = "TP. Hồ Chí Minh"
        elif re.search(r"^lào cai$", raw, re.I):
            t = "Lào Cai"
        else:
            t = canon_tinh(raw)
        if not tinh_moi:
            tinh_moi = t
        parts.append(t)

    # huyện / quận / xã / phường / KCN — lấy làm khóa phụ
    for m in re.finditer(
        r"(?:huyện|quận|thị xã|thành phố|phường|xã|KCN|khu vực)\s+([^,.;()]+)",
        text,
        flags=re.IGNORECASE,
    ):
        frag = normalize_space(m.group(1))
        if frag and len(frag) < 60:
            parts.append(frag)
            if re.match(r"huyện", m.group(0), re.I) and not huyen_cu:
                huyen_cu = frag

    # Biên Hòa, Đồng Nai — cặp thành phố + tỉnh
    m_cap = re.search(
        r"(Biên Hòa|Việt Trì|Móng Cái|Từ Sơn|Duy Tiên|Phú Mỹ|Tân Uyên)[,\s]+([^,.;()]+)",
        text,
    )
    if m_cap:
        parts.append(m_cap.group(1))
        t = canon_tinh(m_cap.group(2))
        # bỏ "(địa chỉ...)" đã cắt bởi [^,.;()]
        if re.search(r"đồng nai|hải phòng|bắc ninh|phú thọ|quảng ninh|ninh bình", t, re.I):
            if not tinh_moi or "nay thuộc" not in text.lower():
                # giữ tỉnh trong cặp nếu chưa có tinh_moi từ "nay thuộc"
                if not tinh_moi:
                    tinh_moi = t
            if not tinh_cu and "nay thuộc" in text.lower():
                tinh_cu = t
            parts.append(t)

    # Hải Phòng / Hà Nội đứng độc lập
    for prov in (
        "Hải Phòng",
        "Hà Nội",
        "Đồng Nai",
        "Bắc Ninh",
        "Quảng Ninh",
        "Ninh Bình",
        "Tây Ninh",
        "Khánh Hòa",
        "Lâm Đồng",
        "An Giang",
        "Cà Mau",
        "Đồng Tháp",
        "Cần Thơ",
        "Hưng Yên",
        "Thanh Hóa",
        "Nghệ An",
        "Hà Tĩnh",
        "Quảng Trị",
        "Huế",
        "Đà Nẵng",
        "Gia Lai",
        "Đắk Lắk",
        "Lào Cai",
        "Phú Thọ",
        "Thái Nguyên",
        "Lạng Sơn",
        "Cao Bằng",
        "Lai Châu",
        "Tuyên Quang",
        "Hà Nam",
        "Bình Dương",
        "Bà Rịa – Vũng Tàu",
        "TP. Hồ Chí Minh",
    ):
        if re.search(re.escape(prov), text, re.I) or (
            prov == "Bà Rịa – Vũng Tàu"
            and re.search(r"Bà Rịa.?Vũng Tàu", text, re.I)
        ):
            parts.append(prov)
            if not tinh_moi:
                tinh_moi = prov

    # Nếu có "X cũ" trong ngoặc tên và tinh_moi khác → tinh_cu
    if not tinh_cu:
        m_cu = re.search(r"\(([^)]+?)\s+cũ\)", text, re.I)
        if m_cu:
            tinh_cu = canon_tinh(m_cu.group(1))

    # khoa_dia_ban: gộp duy nhất, bỏ rỗng
    seen: set[str] = set()
    khoa: list[str] = []
    for p in [tinh_moi, tinh_cu, huyen_cu, *parts]:
        p = normalize_space(p)
        if not p:
            continue
        key = p.lower()
        if key in seen:
            continue
        seen.add(key)
        khoa.append(p)

    return tinh_moi, tinh_cu, huyen_cu, "; ".join(khoa)


def extract_so_qd(so_qd_cell: str, phap_ly: str, tinh_trang: str) -> str:
    """Ưu tiên số QĐ cụ thể; bỏ câu 'Chưa xác minh...'."""
    candidates: list[str] = []
    for blob in (so_qd_cell, tinh_trang, phap_ly):
        if not blob:
            continue
        if blob.lower().startswith("chưa xác minh"):
            # vẫn lấy căn cứ định hướng nếu có
            for m in re.finditer(
                r"(\d+[\w./–-]*\s*/\s*QĐ[-\w]*|\d+/NQ-HĐND|QĐ\s*\d+/QĐ[-\w]*)",
                blob,
                flags=re.IGNORECASE,
            ):
                candidates.append(normalize_space(m.group(1)))
            continue
        # số kiểu 1549/QĐ-BGTVT
        for m in re.finditer(
            r"\b\d+[\w./–-]*\s*/\s*(?:QĐ|NQ)[-\w.]*\b",
            blob,
            flags=re.IGNORECASE,
        ):
            candidates.append(normalize_space(m.group(0)))
        if blob and not candidates and "quyết định" not in blob.lower()[:20]:
            if len(blob) < 80:
                candidates.append(blob)
    # unique giữ thứ tự
    out: list[str] = []
    seen: set[str] = set()
    for c in candidates:
        k = c.lower()
        if k not in seen:
            seen.add(k)
            out.append(c)
    return "; ".join(out[:5])


def read_detail_fields(ws: Worksheet) -> dict[str, tuple[str, str]]:
    """hang_muc -> (noi_dung, ô C)."""
    fields: dict[str, tuple[str, str]] = {}
    for r in range(2, (ws.max_row or 1) + 1):
        hm = cell_str(ws.cell(r, 2).value)
        if not hm:
            continue
        val = cell_str(ws.cell(r, 3).value)
        fields[hm] = (val, f"C{r}")
    return fields


def load_tong_hop(ws: Worksheet) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for r in range(4, (ws.max_row or 0) + 1):
        stt = ws.cell(r, 1).value
        if stt is None:
            continue
        rows.append(
            {
                "stt": str(int(stt)),
                "ten": cell_str(ws.cell(r, 2).value),
                "mien": cell_str(ws.cell(r, 3).value),
                "phap_ly": cell_str(ws.cell(r, 4).value),
                "link": cell_str(ws.cell(r, 5).value),
                "o_ten": f"B{r}",
                "o_link": f"E{r}",
            }
        )
    return rows


def build_ten_fields(ten: str) -> tuple[str, str]:
    """ten_chinh + ten_goi_khac (;-separated)."""
    ten = normalize_space(ten)
    stripped = strip_prefix_ten(ten)
    base, alts = split_paren(stripped)
    variants: list[str] = []
    for v in (ten, stripped, base, *alts):
        v = normalize_space(v)
        if v and v not in variants:
            variants.append(v)
        # thêm bản bỏ giai đoạn
        v2 = strip_prefix_ten(v)
        if v2 and v2 not in variants:
            variants.append(v2)
    ten_chinh = variants[0] if variants else ten
    # ten_goi_khac = phần còn lại (không trùng ten_chinh)
    others = [v for v in variants[1:] if v.lower() != ten_chinh.lower()]
    return ten_chinh, "; ".join(others)


def main() -> int:
    if not EXCEL.exists():
        print(f"Không thấy file bàn giao: {EXCEL}", file=sys.stderr)
        return 1

    wb = load_workbook(EXCEL, read_only=False, data_only=True)
    if "Tổng hợp" not in wb.sheetnames:
        print("Thiếu sheet Tổng hợp", file=sys.stderr)
        return 1

    tong = load_tong_hop(wb["Tổng hợp"])
    tong_by_stt = {int(r["stt"]): r for r in tong}

    hien_trang: list[dict[str, str]] = []
    keys: list[dict[str, str]] = []

    # --- Tổng hợp: 1 dòng / ICD ---
    for r in tong:
        hien_trang.append(
            {
                "loai": "tong_hop",
                "stt": r["stt"],
                "ten": r["ten"],
                "sheet": "Tổng hợp",
                "hang_muc": "Tổng hợp",
                "noi_dung": r["ten"],
                "o": r["o_ten"],
                "mien": r["mien"],
                "phap_ly": r["phap_ly"],
                "link": r["link"],
            }
        )

    # --- Chi tiết: 13 hạng mục / sheet ---
    detail_by_stt: dict[int, dict[str, tuple[str, str]]] = {}
    sheet_by_stt: dict[int, str] = {}

    for name in wb.sheetnames:
        if name == "Tổng hợp":
            continue
        stt = sheet_stt(name)
        if stt is None:
            print(f"Cảnh báo: bỏ sheet không có STT: {name!r}", file=sys.stderr)
            continue
        ws = wb[name]
        fields = read_detail_fields(ws)
        detail_by_stt[stt] = fields
        sheet_by_stt[stt] = name
        ten_sheet = fields.get("Tên", ("", ""))[0] or tong_by_stt.get(stt, {}).get("ten", "")
        for hm, (val, o) in fields.items():
            hien_trang.append(
                {
                    "loai": "chi_tiet",
                    "stt": str(stt),
                    "ten": ten_sheet,
                    "sheet": name,
                    "hang_muc": hm,
                    "noi_dung": val,
                    "o": o,
                    "mien": tong_by_stt.get(stt, {}).get("mien", ""),
                    "phap_ly": tong_by_stt.get(stt, {}).get("phap_ly", ""),
                    "link": tong_by_stt.get(stt, {}).get("link", ""),
                }
            )

    # --- Keys: 128 dòng ---
    for stt in range(1, 129):
        th = tong_by_stt[stt]
        fields = detail_by_stt.get(stt)
        sheet_thieu = "0"

        if fields is None:
            # STT 8 — chưa có sheet chi tiết
            fb = dict(STT8_FALLBACK)
            ten = th["ten"] or fb["ten_chinh"]
            ten_chinh, ten_goi = build_ten_fields(ten)
            # gộp fallback tên gọi khác
            extras = [x.strip() for x in fb["ten_goi_khac"].split(";") if x.strip()]
            existing = {ten_chinh.lower(), *(x.strip().lower() for x in ten_goi.split(";") if x.strip())}
            for e in extras:
                if e.lower() not in existing:
                    ten_goi = f"{ten_goi}; {e}" if ten_goi else e
                    existing.add(e.lower())
            tinh_moi = fb["tinh_moi"]
            tinh_cu = fb["tinh_cu"]
            huyen_cu = fb["huyen_cu"]
            khoa = fb["khoa_dia_ban"]
            so_qd = extract_so_qd("", th["phap_ly"], "")
            chu = fb["chu_dau_tu"]
            dv = fb["don_vi_khai_thac"]
            vi_tri = fb["vi_tri"]
            sheet_name = fb["sheet"]
            sheet_thieu = "1"
        else:
            ten = fields.get("Tên", ("", ""))[0] or th["ten"]
            ten_chinh, ten_goi = build_ten_fields(ten)
            vi_tri = fields.get("Vị trí", ("", ""))[0]
            tinh_moi, tinh_cu, huyen_cu, khoa = parse_vi_tri(vi_tri)
            # bổ sung địa bàn từ ngoặc trong tên (Hưng Yên, An Giang, Gò Dầu...)
            place_hints = {
                "hưng yên",
                "long an",
                "an giang",
                "gò dầu",
                "cái mép",
                "tuy phước",
                "lộc ninh",
                "hậu giang",
                "tp.hcm",
                "bình dương",
            }
            _, alts = split_paren(strip_prefix_ten(ten))
            for a in alts:
                a_clean = re.sub(r"\s+cũ\s*$", "", a, flags=re.I).strip()
                if a_clean and a_clean.lower() not in khoa.lower():
                    khoa = f"{khoa}; {a_clean}" if khoa else a_clean
                if re.search(r"\bcũ\b", a, re.I) and not tinh_cu:
                    tinh_cu = canon_tinh(a_clean)
                elif a_clean.lower() in place_hints:
                    if not tinh_cu and a_clean.lower() != (tinh_moi or "").lower():
                        # Long An trong tên khi tỉnh mới = Tây Ninh
                        if a_clean.lower() != (tinh_moi or "").lower():
                            tinh_cu = tinh_cu or canon_tinh(a_clean)
                    if not huyen_cu and a_clean.lower() in {"gò dầu", "cái mép", "tuy phước", "lộc ninh"}:
                        huyen_cu = a_clean
            so_qd = extract_so_qd(
                fields.get("Số QĐ thành lập", ("", ""))[0],
                th["phap_ly"],
                fields.get("Tình trạng", ("", ""))[0],
            )
            chu = fields.get("Chủ đầu tư", ("", ""))[0]
            dv = fields.get("Đơn vị khai thác", ("", ""))[0]
            sheet_name = sheet_by_stt.get(stt, "")

        # luôn gắn tên ngắn vào khóa địa bàn (giúp khớp + phân biệt R7)
        short, _ = split_paren(strip_prefix_ten(ten_chinh))
        if short and short.lower() not in khoa.lower():
            khoa = f"{khoa}; {short}" if khoa else short

        # đảm bảo có ít nhất một khóa địa bàn
        if not khoa:
            if huyen_cu:
                khoa = huyen_cu
            elif tinh_moi:
                khoa = tinh_moi
            elif tinh_cu:
                khoa = tinh_cu
            else:
                khoa = short or ten_chinh

        keys.append(
            {
                "stt": str(stt),
                "ten_chinh": ten_chinh,
                "ten_goi_khac": ten_goi,
                "tinh_moi": tinh_moi,
                "tinh_cu": tinh_cu,
                "huyen_cu": huyen_cu,
                "khoa_dia_ban": khoa,
                "so_qd": so_qd,
                "chu_dau_tu": chu,
                "don_vi_khai_thac": dv,
                "mien": th["mien"],
                "phap_ly": th["phap_ly"],
                "link": th["link"],
                "vi_tri": vi_tri,
                "sheet": sheet_name,
                "sheet_thieu": sheet_thieu,
                "nhom_trung": NHOM_TRUNG.get(stt, ""),
            }
        )

    wb.close()

    OUT_HIEN_TRANG.parent.mkdir(parents=True, exist_ok=True)

    ht_fields = [
        "loai",
        "stt",
        "ten",
        "sheet",
        "hang_muc",
        "noi_dung",
        "o",
        "mien",
        "phap_ly",
        "link",
    ]
    with OUT_HIEN_TRANG.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=ht_fields)
        w.writeheader()
        w.writerows(hien_trang)

    key_fields = [
        "stt",
        "ten_chinh",
        "ten_goi_khac",
        "tinh_moi",
        "tinh_cu",
        "huyen_cu",
        "khoa_dia_ban",
        "so_qd",
        "chu_dau_tu",
        "don_vi_khai_thac",
        "mien",
        "phap_ly",
        "link",
        "vi_tri",
        "sheet",
        "sheet_thieu",
        "nhom_trung",
    ]
    with OUT_KEYS.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=key_fields)
        w.writeheader()
        w.writerows(keys)

    # --- Báo cáo ---
    n_th = sum(1 for r in hien_trang if r["loai"] == "tong_hop")
    n_ct = sum(1 for r in hien_trang if r["loai"] == "chi_tiet")
    n_sheets = len(detail_by_stt)
    per_sheet = {}
    for r in hien_trang:
        if r["loai"] != "chi_tiet":
            continue
        per_sheet[r["stt"]] = per_sheet.get(r["stt"], 0) + 1
    bad_hm = {s: n for s, n in per_sheet.items() if n != 13}

    thieu_khoa = [k for k in keys if not normalize_space(k["khoa_dia_ban"])]
    nhom: dict[str, list[str]] = {}
    for k in keys:
        if k["nhom_trung"]:
            nhom.setdefault(k["nhom_trung"], []).append(k["stt"])

    print("=== P0.3 icd_hien_trang.csv ===")
    print(f"  File: {OUT_HIEN_TRANG}")
    print(f"  Tổng dòng dữ liệu: {len(hien_trang)} (không kể header)")
    print(f"  tong_hop: {n_th} (cần 128) → {'OK' if n_th == 128 else 'FAIL'}")
    print(
        f"  chi_tiet: {n_ct} từ {n_sheets} sheet"
        f" (cần 127×13={127*13}) → {'OK' if n_ct == 127 * 13 and n_sheets == 127 else 'FAIL'}"
    )
    if bad_hm:
        print(f"  Sheet lệch 13 hạng mục: {bad_hm}")
    missing_detail = [i for i in range(1, 129) if i not in detail_by_stt]
    print(f"  Sheet chi tiết thiếu (dự kiến [8]): {missing_detail}")

    print("=== P0.4 icd_keys.csv ===")
    print(f"  File: {OUT_KEYS}")
    print(f"  Số dòng: {len(keys)} (cần 128) → {'OK' if len(keys) == 128 else 'FAIL'}")
    print(f"  ICD thiếu khóa địa bàn: {len(thieu_khoa)}")
    for k in thieu_khoa:
        print(f"    - STT {k['stt']}: {k['ten_chinh']}")
    print("  Nhóm trùng tên (R7):")
    for name, stts in nhom.items():
        print(f"    - {name}: STT {', '.join(stts)}")

    ok_p03 = n_th == 128 and n_ct == 127 * 13 and n_sheets == 127
    ok_p04 = len(keys) == 128 and not thieu_khoa
    print("=== Kết luận ===")
    print(f"  P0.3: {'ĐẠT' if ok_p03 else 'CHƯA ĐẠT'}")
    print(f"  P0.4: {'ĐẠT' if ok_p04 else 'CHƯA ĐẠT'}")
    return 0 if ok_p03 and ok_p04 else 2


if __name__ == "__main__":
    raise SystemExit(main())
